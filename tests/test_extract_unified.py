"""Tests for F-5: 统一提取接口与格式自动检测."""

import io

import docx
import pytest
from fpdf import FPDF
from httpx import AsyncClient


def _make_docx_bytes(texts: list[str]) -> bytes:
    doc = docx.Document()
    for t in texts:
        doc.add_paragraph(t)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_pdf_bytes(texts: list[str]) -> bytes:
    pdf = FPDF()
    pdf.set_font("Helvetica", size=12)
    for t in texts:
        pdf.add_page()
        pdf.cell(text=t)
    buf = io.BytesIO()
    pdf.output(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_unified_extract_text(client: AsyncClient):
    resp = await client.post(
        "/extract",
        files={"file": ("doc.txt", b"hello world", "text/plain")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["format"] == "text"
    assert "hello world" in data["text"]


@pytest.mark.asyncio
async def test_unified_extract_pdf(client: AsyncClient):
    pdf_bytes = _make_pdf_bytes(["PDF content"])
    resp = await client.post(
        "/extract",
        files={"file": ("doc.pdf", pdf_bytes, "application/pdf")},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["format"] == "pdf"
    assert "PDF content" in data["text"]
    assert data["page_count"] == 1


@pytest.mark.asyncio
async def test_unified_extract_docx(client: AsyncClient):
    docx_bytes = _make_docx_bytes(["Word content"])
    resp = await client.post(
        "/extract",
        files={
            "file": (
                "doc.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["format"] == "docx"
    assert "Word content" in data["text"]


@pytest.mark.asyncio
async def test_unified_unknown_format_returns_422(client: AsyncClient):
    resp = await client.post(
        "/extract",
        files={"file": ("data.xyz", b"???", "application/octet-stream")},
    )
    assert resp.status_code == 422
    assert ".txt" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_unified_no_extension_uses_mime(client: AsyncClient):
    """File with no extension but text/plain MIME should be detected as text."""
    resp = await client.post(
        "/extract",
        files={"file": ("noext", b"content", "text/plain")},
    )
    assert resp.status_code == 200
    assert resp.json()["format"] == "text"


@pytest.mark.asyncio
async def test_unified_response_has_required_fields(client: AsyncClient):
    """Response must contain format, text, page_count, char_count, extraction_time_ms."""
    resp = await client.post(
        "/extract",
        files={"file": ("doc.txt", b"test", "text/plain")},
    )
    data = resp.json()
    for key in ("format", "text", "page_count", "char_count", "extraction_time_ms"):
        assert key in data


@pytest.mark.asyncio
async def test_char_count_matches_text_length(client: AsyncClient):
    resp = await client.post(
        "/extract",
        files={"file": ("doc.txt", b"hello", "text/plain")},
    )
    data = resp.json()
    assert data["char_count"] == len(data["text"])


@pytest.mark.asyncio
async def test_extraction_time_ms_is_positive(client: AsyncClient):
    resp = await client.post(
        "/extract",
        files={"file": ("doc.txt", b"test", "text/plain")},
    )
    assert resp.json()["extraction_time_ms"] >= 0


@pytest.mark.asyncio
async def test_doc_legacy_format_rejected_with_hint(client: AsyncClient):
    """Uploading a .doc file should return 422 with legacy hint."""
    resp = await client.post(
        "/extract",
        files={"file": ("legacy.doc", b"fake ole data", "application/octet-stream")},
    )
    assert resp.status_code == 422
    assert "Legacy .doc" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_docx_page_count_is_null(client: AsyncClient):
    """Word documents should have page_count = null (not computed)."""
    docx_bytes = _make_docx_bytes(["Content"])
    resp = await client.post(
        "/extract",
        files={
            "file": (
                "doc.docx",
                docx_bytes,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert resp.status_code == 200
    assert resp.json()["page_count"] is None
